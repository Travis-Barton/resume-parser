"""Acceptance checks for the existing Word templates with synthetic data only."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

from document_renderer import TEMPLATE_DIR, content_markup, render_resume_docx, replace_text_in_paragraph


def sections(**overrides):
    value = {
        "name": "Synthetic Alex Rivera",
        "summary": "Software engineer working with data systems.",
        "skills_and_tech": "Python\nSQL",
        "professional_experience": "Example Co, Remote\t2020–2024\nSoftware Engineer\nBuilt data pipelines\nReviewed releases\n\nSample Lab\t2018–2020\nDeveloper\nMaintained internal tools",
        "education": "BSc Computer Science, Example University, 2018",
    }
    value.update(overrides)
    return value


class DocumentFormattingTests(unittest.TestCase):
    def test_openable_docx_has_name_sections_and_order(self):
        data = render_resume_docx(sections())
        self.assertIsInstance(data, bytes)
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Synthetic Alex Rivera", text)
        self.assertNotIn("APPLICANT_NAME", text)
        headers = ["SUMMARY", "SKILLS AND TECHNOLOGIES", "PROFESSIONAL EXPERIENCE", "EDUCATION"]
        positions = [text.index(header) for header in headers]
        self.assertEqual(positions, sorted(positions))
        for expected in ("Python", "Software Engineer", "Built data pipelines", "BSc Computer Science"):
            self.assertIn(expected, text)

    def test_existing_styles_and_bullets_preserved(self):
        doc = Document(BytesIO(render_resume_docx(sections())))
        by_text = {p.text: p for p in doc.paragraphs}
        self.assertEqual(by_text["Python"].style.name, "List Bullet 2")
        self.assertEqual(by_text["Built data pipelines"].style.name, "List Bullet 2")
        self.assertEqual(by_text["Software Engineer"].style.name, "Travis Normal Bold 2")
        self.assertEqual(by_text["Example Co, Remote\t2020–2024"].style.name, "Travis Normal Bold 2")
        self.assertTrue(all(run.font.name == "Calibri" for p in doc.paragraphs for run in p.runs))
        self.assertNotIn("•", by_text["Python"].text)

    def test_empty_job_headers_do_not_turn_bullets_into_headings(self):
        value = sections(professional_experience="\t\n\t\nBuilt a synthetic tool\n\nExample Co\n\t\nTested a synthetic tool")
        doc = Document(BytesIO(render_resume_docx(value)))
        by_text = {p.text: p for p in doc.paragraphs}
        self.assertEqual(by_text["Built a synthetic tool"].style.name, "List Bullet 2")
        self.assertEqual(by_text["Tested a synthetic tool"].style.name, "List Bullet 2")
        self.assertEqual(by_text["Example Co"].style.name, "Travis Normal Bold 2")

    def test_empty_sections_add_no_invented_information(self):
        value = {key: "" for key in sections()}
        doc = Document(BytesIO(render_resume_docx(value)))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("APPLICANT_NAME", text)
        self.assertNotIn("Name Not Found", text)
        self.assertFalse(any(p.style.name == "List Bullet 2" for p in doc.paragraphs if p.text))

    def test_domain_words_and_punctuation_are_not_destructively_cleaned(self):
        value = sections(summary="Created education tools; skills_and_tech is a source field. Uses C++ and .NET.")
        doc = Document(BytesIO(render_resume_docx(value)))
        self.assertIn(value["summary"], [p.text for p in doc.paragraphs])

    def test_legacy_section_headers_fences_and_list_markers_are_cleaned(self):
        value = sections(
            summary="```summary\nSUMMARY\nSoftware engineer.\n```",
            skills_and_tech="``` skills_and_tech\nSKILLS AND TECHNOLOGIES\n• Python\n- SQL\n* C++\n```",
            professional_experience=(
                "```professional_experience\nPROFESSIONAL EXPERIENCE\n"
                "Example Co, Remote\t2020-2024\nSoftware Engineer\n• Built data pipelines\n- Reviewed releases\n```"
            ),
            education="```education\neducation\nBSc Computer Science, Example University\n```",
        )
        doc = Document(BytesIO(render_resume_docx(value)))
        texts = [p.text for p in doc.paragraphs]
        joined = "\n".join(texts)
        self.assertNotIn("```", joined)
        for header in ("SUMMARY", "SKILLS AND TECHNOLOGIES", "PROFESSIONAL EXPERIENCE", "EDUCATION"):
            self.assertEqual(sum(p.strip() == header for p in texts), 1)
        by_text = {p.text: p for p in doc.paragraphs}
        for bullet in ("Python", "SQL", "C++", "Built data pipelines", "Reviewed releases"):
            self.assertEqual(by_text[bullet].style.name, "List Bullet 2")

    def test_experience_cleanup_stops_at_leaked_section_heading(self):
        content = "Example Co\nEngineer\nBuilt tools\n\nEducation\nSHOULD_NOT_EXPORT"
        self.assertEqual(content_markup(content, "professional_experience"), "Example Co\nEngineer\nBuilt tools")

    def test_cleanup_preserves_domain_words_and_tabs_in_job_header(self):
        content = "Education Labs, Remote\t2020-2024\nEngineer\nBuilt Skills Summary reporting"
        self.assertEqual(content_markup(content, "professional_experience"), content)
        self.assertEqual(content_markup("*nix\n.NET\nC++", "skills_and_tech"), "*nix\n.NET\nC++")

    def test_template_tokens_split_across_runs_keep_neighbor_text(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Prefix ").bold = True
        p.add_run("APPLI")
        p.add_run("CANT_NAME")
        p.add_run(" suffix")
        replace_text_in_paragraph(p, "APPLICANT_NAME", "Synthetic Casey")
        self.assertEqual(p.text, "Prefix Synthetic Casey suffix")
        self.assertTrue(p.runs[0].bold)

    def test_multiple_template_tokens_do_not_reprocess_replacement(self):
        p = Document().add_paragraph("APPLICANT_NAME / APPLICANT_NAME")
        replace_text_in_paragraph(p, "APPLICANT_NAME", "APPLICANT_NAME Junior")
        self.assertEqual(p.text, "APPLICANT_NAME Junior / APPLICANT_NAME Junior")

    def test_rendering_does_not_write_shared_output_or_change_templates(self):
        root = Path(__file__).resolve().parents[1]
        before_names = set(root.iterdir())
        before_templates = {p.name: sha256(p.read_bytes()).hexdigest() for p in TEMPLATE_DIR.glob("*.docx")}
        first = Document(BytesIO(render_resume_docx(sections(name="Synthetic One"))))
        second = Document(BytesIO(render_resume_docx(sections(name="Synthetic Two"))))
        self.assertIn("Synthetic One", "\n".join(p.text for p in first.paragraphs))
        self.assertNotIn("Synthetic One", "\n".join(p.text for p in second.paragraphs))
        self.assertEqual(before_names, set(root.iterdir()))
        self.assertEqual(before_templates, {p.name: sha256(p.read_bytes()).hexdigest() for p in TEMPLATE_DIR.glob("*.docx")})

    def test_invalid_section_contract_is_rejected(self):
        for value in (sections(extra="not allowed"), {"name": "Synthetic"}, sections(summary=None)):
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    render_resume_docx(value)


if __name__ == "__main__":
    unittest.main()
