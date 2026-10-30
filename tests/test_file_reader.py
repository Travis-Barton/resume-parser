"""File-reading checks use only synthetic in-memory applicant content."""

from io import BytesIO, StringIO
import unittest
from unittest.mock import patch

from docx import Document

from utils import MAX_UPLOAD_BYTES, ResumeFileError, file_reader


class Upload(BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name


def synthetic_pdf(text):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii"))
    return bytes(data)


class FileReaderTests(unittest.TestCase):
    def test_utf8_text_and_bom_are_unicode(self):
        expected = "Synthetic Renée — Python"
        self.assertEqual(file_reader(Upload(expected.encode("utf-8-sig"), "RESUME.TXT")), expected)

    def test_file_cursor_is_reset(self):
        uploaded = Upload(b"Synthetic resume", "resume.txt")
        uploaded.read()
        self.assertEqual(file_reader(uploaded), "Synthetic resume")

    def test_existing_text_stream_is_supported(self):
        uploaded = StringIO("Synthetic text stream")
        uploaded.name = "resume.txt"
        self.assertEqual(file_reader(uploaded), "Synthetic text stream")

    def test_non_utf8_is_rejected_without_bytes_in_message(self):
        with self.assertRaisesRegex(ResumeFileError, "not UTF-8"):
            file_reader(Upload(b"\xff\xfeinvalid", "resume.txt"))

    def test_legacy_doc_is_explicitly_rejected(self):
        with self.assertRaisesRegex(ResumeFileError, "Legacy .doc"):
            file_reader(Upload(b"not a docx", "RESUME.DOC"))

    def test_unknown_type_is_rejected(self):
        with self.assertRaisesRegex(ResumeFileError, "DOCX, PDF"):
            file_reader(Upload(b"synthetic", "resume.html"))

    def test_size_limit(self):
        with self.assertRaisesRegex(ResumeFileError, "too large"):
            file_reader(Upload(b"x" * (MAX_UPLOAD_BYTES + 1), "resume.txt"))

    def test_docx_keeps_table_text_in_document_order(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Synthetic Candidate"
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "SQL"
        doc.add_paragraph("After table")
        data = BytesIO()
        doc.save(data)
        result = file_reader(Upload(data.getvalue(), "resume.DOCX"))
        self.assertIn("Synthetic Candidate", result)
        self.assertLess(result.index("Before table"), result.index("Python\tSQL"))
        self.assertLess(result.index("Python\tSQL"), result.index("After table"))

    def test_text_based_pdf_extracts_without_a_tempfile(self):
        result = file_reader(Upload(synthetic_pdf("Synthetic PDF resume"), "resume.PDF"))
        self.assertIn("Synthetic PDF resume", result)

    def test_parser_exception_is_not_exposed(self):
        with patch("utils.extract_text", side_effect=RuntimeError("SENTINEL_PRIVATE_CONTENT")):
            with self.assertRaises(ResumeFileError) as caught:
                file_reader(Upload(b"synthetic", "resume.pdf"))
        self.assertNotIn("SENTINEL", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
